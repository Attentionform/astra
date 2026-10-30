"""Writes python-docx-three-sections.docx: three sections made by python-docx.

Section 1 is portrait Letter. Section 2 is landscape with its own header.
Section 3 is continuous, portrait, with a different first-page header.
Run: python3 generate_python_docx.py   (needs python-docx 1.2.0)
"""
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Inches

doc = Document()
doc.add_paragraph("Section one text")
doc.add_paragraph("Section one, second paragraph")

second = doc.add_section(WD_SECTION.NEW_PAGE)
second.orientation = WD_ORIENT.LANDSCAPE
second.page_width, second.page_height = Inches(11), Inches(8.5)
second.header.is_linked_to_previous = False
second.header.paragraphs[0].text = "Header of section two"
doc.add_paragraph("Section two text (landscape)")

third = doc.add_section(WD_SECTION.CONTINUOUS)
third.orientation = WD_ORIENT.PORTRAIT
third.page_width, third.page_height = Inches(8.5), Inches(11)
third.different_first_page_header_footer = True
third.first_page_header.is_linked_to_previous = False
third.first_page_header.paragraphs[0].text = "First page header of section three"
doc.add_paragraph("Section three text")

doc.save("python-docx-three-sections.docx")
